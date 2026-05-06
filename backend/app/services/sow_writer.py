"""Stage 10 — generate per-trade-package SOW Word documents.

For each TradePackage in a run, render a `.docx` containing:
    - Header: project name + lifecycle state + package label + CSI divisions
              + trust score + bilateral coverage
    - Per-CSI-section table of scope items (description / qty / unit /
      evidence tier / citation hyperlinks)
    - Citations footer

Citations link back to the backend page-viewer URL, so opening the doc in
Microsoft Word and clicking a citation jumps to the source page.

No LLM calls — pure rendering.
"""

from __future__ import annotations

import logging
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

from docx import Document as Docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from ..config import settings
from ..database import SessionLocal
from ..models import (
    Project,
    ScopeCitation,
    ScopeExtractionRun,
    ScopeItem,
    TradePackage,
)

log = logging.getLogger(__name__)


@dataclass
class GeneratedSOW:
    package_id: str
    package_label: str
    file_path: Path
    file_size: int
    item_count: int


# -----------------------------------------------------------------------------
# python-docx hyperlink helper (the library doesn't ship one)
# -----------------------------------------------------------------------------


def _add_hyperlink(paragraph, url: str, text: str) -> None:
    """Insert a clickable hyperlink with blue underlined formatting."""
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F6FEB")
    rPr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)
    new_run.append(rPr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hl.append(new_run)
    paragraph._p.append(hl)


# -----------------------------------------------------------------------------
# Layout helpers
# -----------------------------------------------------------------------------


def _page_viewer_url(
    base_url: str, project_id: str, document_id: str, page: int
) -> str:
    return (
        f"{base_url}/api/projects/{project_id}/documents/{document_id}"
        f"/pages/{page}/image"
    )


def _evidence_tier_label(tier: str | None) -> str:
    if tier == "EXPLICITLY_CITED":
        return "Explicit"
    if tier == "INFERRED_HIGH_CONFIDENCE":
        return "Inferred / High"
    if tier == "INFERRED_LOW_CONFIDENCE":
        return "Inferred / Low"
    return "—"


# -----------------------------------------------------------------------------
# SOW renderer
# -----------------------------------------------------------------------------


def _render_sow(
    doc: Docx,
    *,
    project: Project,
    run: ScopeExtractionRun,
    pkg: TradePackage,
    items: list[ScopeItem],
    citations_by_item: dict[str, list[ScopeCitation]],
    api_base_url: str,
) -> int:
    # Title
    title = doc.add_heading(pkg.package_label, level=0)
    for run_text in title.runs:
        run_text.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Header metadata
    meta = doc.add_paragraph()
    meta.add_run(f"Project: {project.name}\n").bold = True
    meta.add_run(f"Lifecycle: {project.lifecycle_state}\n")
    if pkg.csi_divisions:
        meta.add_run(f"CSI Divisions: {', '.join(pkg.csi_divisions)}\n")
    meta.add_run(f"Items in package: {len(items)}\n")
    if pkg.bilateral_count is not None and pkg.item_count > 0:
        bilateral_pct = pkg.bilateral_count / pkg.item_count * 100
        meta.add_run(
            f"Bilateral coverage: "
            f"{pkg.bilateral_count}/{pkg.item_count} ({bilateral_pct:.0f}%)\n"
        )
    if pkg.avg_confidence is not None:
        meta.add_run(f"Average extraction confidence: {pkg.avg_confidence:.0%}\n")
    if run.trust_score is not None:
        meta.add_run(
            f"Project trust score: {run.trust_score:.2f} "
            f"({(run.trust_score_components or {}).get('tier', '—')})\n"
        )

    # Group items by csi_section for readable sections
    by_section: dict[str, list[ScopeItem]] = defaultdict(list)
    for it in items:
        by_section[it.csi_code].append(it)

    rendered = 0
    for section_code in sorted(by_section.keys()):
        section_items = by_section[section_code]
        title_text = section_items[0].section_title or "(untitled section)"
        doc.add_heading(f"{section_code} — {title_text}", level=2)

        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Description"
        hdr[1].text = "Quantity"
        hdr[2].text = "Unit"
        hdr[3].text = "Tier"
        hdr[4].text = "Citations"

        for it in section_items:
            if it.verifier_status == "rejected":
                continue
            row = table.add_row().cells
            row[0].text = it.description
            row[1].text = (it.quantity or "—").strip() or "—"
            row[2].text = (it.unit or "").strip() or "—"
            row[3].text = _evidence_tier_label(it.evidence_tier)
            cit_para = row[4].paragraphs[0]
            cits = citations_by_item.get(it.id, [])
            for idx, cit in enumerate(cits[:5]):
                if idx > 0:
                    cit_para.add_run("  ")
                if cit.document_id and cit.page_number:
                    label = (
                        cit.sheet_number
                        or f"p.{cit.page_number}"
                    )
                    _add_hyperlink(
                        cit_para,
                        _page_viewer_url(
                            api_base_url, project.id, cit.document_id, cit.page_number
                        ),
                        f"[{label}]",
                    )
                else:
                    cit_para.add_run("[—]")
            if not cits:
                cit_para.add_run("(none)")
            rendered += 1

    # Footer
    footer = doc.add_paragraph()
    fr = footer.add_run(
        "\nGenerated by ECS Estimating Agent. "
        "Citations link to the source PDF page in the backend viewer."
    )
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    return rendered


# -----------------------------------------------------------------------------
# Public entry
# -----------------------------------------------------------------------------


async def write_sows_for_run(
    run_id: str, *, output_dir: Path, api_base_url: str
) -> list[GeneratedSOW]:
    """Write one .docx per TradePackage in the run. Returns metadata list."""
    output_dir.mkdir(parents=True, exist_ok=True)
    generated: list[GeneratedSOW] = []

    async with SessionLocal() as db:
        run = await db.get(ScopeExtractionRun, run_id)
        if run is None:
            raise ValueError(f"sow_writer: run {run_id} not found")
        project = await db.get(Project, run.project_id)
        if project is None:
            raise ValueError(f"sow_writer: project for run {run_id} missing")

        packages = (
            await db.execute(
                select(TradePackage).where(TradePackage.run_id == run_id)
            )
        ).scalars().all()
        if not packages:
            log.info("sow_writer: no packages for run %s", run_id)
            return []

        for pkg in packages:
            items = (
                await db.execute(
                    select(ScopeItem)
                    .where(ScopeItem.package_id == pkg.id)
                    .order_by(ScopeItem.csi_code)
                )
            ).scalars().all()
            if not items:
                continue

            cit_rows = (
                await db.execute(
                    select(ScopeCitation).where(
                        ScopeCitation.scope_item_id.in_([i.id for i in items])
                    )
                )
            ).scalars().all()
            citations_by_item: dict[str, list[ScopeCitation]] = defaultdict(list)
            for c in cit_rows:
                citations_by_item[c.scope_item_id].append(c)

            doc = Docx()
            count = _render_sow(
                doc,
                project=project,
                run=run,
                pkg=pkg,
                items=items,
                citations_by_item=dict(citations_by_item),
                api_base_url=api_base_url,
            )
            file_path = output_dir / f"{pkg.package_key}.docx"
            doc.save(file_path)
            generated.append(
                GeneratedSOW(
                    package_id=pkg.id,
                    package_label=pkg.package_label,
                    file_path=file_path,
                    file_size=file_path.stat().st_size,
                    item_count=count,
                )
            )

    log.info(
        "sow_writer: wrote %d SOWs to %s for run %s",
        len(generated),
        output_dir,
        run_id,
    )
    return generated


def default_api_base_url() -> str:
    """Best-effort URL for hyperlink targets. Override via env later."""
    return f"http://localhost:{settings.app_port}"
